from __future__ import annotations

from csv import DictReader
import json
from pathlib import Path
from typing import Iterable

from ..types import ClipKey, DirectionalCounts, UpstreamDirection, WeakCountRecord


def _load_tabular_rows(path: Path) -> list[dict[str, object]]:
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return list(DictReader(handle, delimiter=delimiter))
    if suffix == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(payload, list):
            raise ValueError("JSON weak-label files must contain a list of objects")
        return payload
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook  # type: ignore
        except ImportError as exc:
            raise RuntimeError("openpyxl is required to read XLSX weak-label files") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        header = [str(cell) for cell in rows[0]]
        return [dict(zip(header, row, strict=False)) for row in rows[1:]]
    raise ValueError(f"Unsupported weak-label file format: {path.suffix}")


def _to_int(value: object | None) -> int | None:
    if value is None or value == "":
        return None
    return int(float(value))


def _row_to_record(row: dict[str, object], source_path: Path) -> WeakCountRecord:
    upstream_direction = UpstreamDirection.from_optional(str(row.get("upstream_direction") or "").strip() or None)
    left_value = _to_int(row.get("left_count"))
    right_value = _to_int(row.get("right_count"))
    if left_value is not None or right_value is not None:
        counts = DirectionalCounts.from_left_right(left_value, right_value)
    else:
        counts = DirectionalCounts.from_upstream_downstream(
            _to_int(row.get("upstream_count")),
            _to_int(row.get("downstream_count")),
            upstream_direction,
        )
    return WeakCountRecord(
        domain=str(row["domain"]),
        clip_id=str(row["clip_id"]),
        counts=counts,
        source_type="weak",
        upstream_direction=upstream_direction,
        source_path=source_path,
    )


def load_weak_count_records(path: str | Path) -> list[WeakCountRecord]:
    source_path = Path(path)
    return [_row_to_record(row, source_path) for row in _load_tabular_rows(source_path)]


def index_weak_counts(records: Iterable[WeakCountRecord]) -> dict[ClipKey, WeakCountRecord]:
    return {record.key: record for record in records}

